"""Integration tests for Articles API endpoints."""

import io

from openpyxl import Workbook


class TestArticlesRetrieval:
    """Tests for articles retrieval endpoints."""

    def test_get_all_articles_requires_auth(self, client):
        """Test that getting all articles requires authentication."""
        response = client.get("/api/articles/all")

        assert response.status_code == 401

    def test_get_all_articles_with_auth(self, client, auth_headers, test_article):
        """Test getting all articles with authentication."""
        response = client.get("/api/articles/all", headers=auth_headers)

        assert response.status_code == 200
        data = response.get_json()
        assert "data" in data

    def test_get_all_articles_pagination(self, client, auth_headers, test_article):
        """Test pagination parameters."""
        response = client.get(
            "/api/articles/all?page=1&per_page=10", headers=auth_headers
        )

        assert response.status_code == 200
        data = response.get_json()
        assert "data" in data

    def test_get_article_info_requires_auth(self, client, test_article):
        """Test that getting article info requires authentication."""
        response = client.get(f"/api/articles/info/{test_article.id}")

        assert response.status_code == 401

    def test_get_article_info_with_auth(self, client, auth_headers, test_article):
        """Test getting specific article info with authentication."""
        response = client.get(
            f"/api/articles/info/{test_article.id}", headers=auth_headers
        )

        assert response.status_code == 200
        data = response.get_json()
        assert data["id"] == test_article.id
        assert data["nom"] == test_article.nom

    def test_get_nonexistent_article(self, client, auth_headers):
        """Test getting non-existent article."""
        response = client.get("/api/articles/info/99999", headers=auth_headers)

        assert response.status_code == 404


class TestArticlesCreate:
    """Tests for creating articles."""

    def test_create_article_requires_auth(self, client):
        """Test that creating an article requires authentication."""
        response = client.post(
            "/api/articles/create",
            json={
                "nom": "Test Article",
                "designation": "TEST-001",
                "reference": "REF-001",
                "prix_achat_HT": 100.0,
                "taux_tva": 20.0,
            },
        )

        assert response.status_code == 401

    def test_create_article_success(
        self, client, auth_headers, taux_tva_20, test_parameters
    ):
        """Test successful article creation."""
        response = client.post(
            "/api/articles/create",
            json={
                "nom": "New Article",
                "designation": "NEW-001",
                "reference": "REF-NEW-001",
                "prix_achat_HT": 100.0,
                "taux_tva": 20.0,
            },
            headers=auth_headers,
        )

        assert response.status_code == 200
        data = response.get_json()
        assert "id" in data

    def test_create_article_without_tva(self, client, auth_headers):
        """Test creating article without TVA."""
        response = client.post(
            "/api/articles/create",
            json={
                "nom": "Article sans TVA",
                "designation": "NO-TVA-001",
                "reference": "REF-NO-TVA-001",
                "prix_achat_HT": 100.0,
                "taux_tva": None,
            },
            headers=auth_headers,
        )

        assert response.status_code == 400
        data = response.get_json()
        assert "error" in data

    def test_create_article_with_invalid_tva_format(self, client, auth_headers):
        """Test creating article with invalid TVA format."""
        response = client.post(
            "/api/articles/create",
            json={
                "nom": "Article TVA invalide",
                "designation": "TVA-INV-001",
                "reference": "REF-TVA-INV-001",
                "prix_achat_HT": 100.0,
                "taux_tva": "not-a-number",
            },
            headers=auth_headers,
        )

        assert response.status_code == 400
        assert "invalide" in response.get_json()["error"]

    def test_create_article_with_unknown_tva_rate(self, client, auth_headers):
        """Test creating article with unknown TVA rate."""
        response = client.post(
            "/api/articles/create",
            json={
                "nom": "Article TVA inconnue",
                "designation": "TVA-UNK-001",
                "reference": "REF-TVA-UNK-001",
                "prix_achat_HT": 100.0,
                "taux_tva": 5,
            },
            headers=auth_headers,
        )

        assert response.status_code == 400
        assert "introuvable" in response.get_json()["error"]


class TestArticlesUpdate:
    """Tests for updating articles."""

    def test_update_article_requires_auth(self, client, test_article):
        """Test that updating an article requires authentication."""
        response = client.post(
            f"/api/articles/update/{test_article.id}",
            json={
                "nom": "Updated Article",
                "designation": "UPD-001",
                "reference": "REF-UPD-001",
                "prix_achat_HT": 150.0,
                "taux_tva": 20.0,
            },
        )

        assert response.status_code == 401

    def test_update_article_success(
        self, client, auth_headers, test_article, taux_tva_20, test_parameters
    ):
        """Test successful article update."""
        response = client.post(
            f"/api/articles/update/{test_article.id}",
            json={
                "nom": "Updated Article Name",
                "designation": "UPD-REF",
                "reference": "REF-UPD-REF",
                "prix_achat_HT": 200.0,
                "taux_tva": 20.0,
            },
            headers=auth_headers,
        )

        assert response.status_code == 200
        data = response.get_json()
        assert "id" in data
        assert data["id"] == test_article.id

    def test_update_nonexistent_article(self, client, auth_headers):
        """Test updating non-existent article."""
        response = client.post(
            "/api/articles/update/99999",
            json={
                "nom": "Test",
                "designation": "TEST",
                "reference": "REF-TEST",
                "prix_achat_HT": 100.0,
                "taux_tva": 20.0,
            },
            headers=auth_headers,
        )

        assert response.status_code == 404

    def test_update_article_without_tva(self, client, auth_headers, test_article):
        """Test updating article without TVA."""
        response = client.post(
            f"/api/articles/update/{test_article.id}",
            json={
                "nom": "Article sans TVA",
                "designation": "NO-TVA",
                "reference": "REF-NO-TVA",
                "prix_achat_HT": 100.0,
                "taux_tva": None,
            },
            headers=auth_headers,
        )

        assert response.status_code == 400

    def test_update_article_with_invalid_tva_format(
        self, client, auth_headers, test_article
    ):
        """Test updating article with invalid TVA format."""
        response = client.post(
            f"/api/articles/update/{test_article.id}",
            json={
                "nom": "Article TVA invalide",
                "designation": "UPD-TVA-INV",
                "reference": "REF-UPD-TVA-INV",
                "prix_achat_HT": 100.0,
                "taux_tva": "abc",
            },
            headers=auth_headers,
        )

        assert response.status_code == 400
        assert "invalide" in response.get_json()["error"]

    def test_create_article_with_location_price(
        self, client, auth_headers, taux_tva_20, test_parameters
    ):
        """Test creating article with location_price (subscription price)."""
        response = client.post(
            "/api/articles/create",
            json={
                "nom": "Article avec Prix Location",
                "designation": "LOCATION-001",
                "reference": "REF-LOCATION-001",
                "prix_achat_HT": 100.0,
                "location_price": 120.0,
                "taux_tva": 20.0,
            },
            headers=auth_headers,
        )

        assert response.status_code == 200
        data = response.get_json()
        assert "id" in data

        # Verify location_price was saved
        response = client.get(f"/api/articles/info/{data['id']}", headers=auth_headers)
        assert response.status_code == 200
        article_data = response.get_json()
        assert article_data["location_price"] == 120.0

    def test_create_article_without_location_price_defaults_to_zero(
        self, client, auth_headers, taux_tva_20, test_parameters
    ):
        """Test creating article without location_price defaults to 0."""
        response = client.post(
            "/api/articles/create",
            json={
                "nom": "Article sans Prix Location",
                "designation": "NO-LOCATION-001",
                "reference": "REF-NO-LOCATION-001",
                "prix_achat_HT": 100.0,
                "taux_tva": 20.0,
            },
            headers=auth_headers,
        )

        assert response.status_code == 200
        data = response.get_json()
        assert "id" in data

        # Verify location_price defaults to 0 when omitted
        response = client.get(f"/api/articles/info/{data['id']}", headers=auth_headers)
        assert response.status_code == 200
        article_data = response.get_json()
        assert article_data["location_price"] == 0.0

    def test_update_article_with_location_price(
        self, client, auth_headers, test_article, taux_tva_20, test_parameters
    ):
        """Test updating article with location_price."""
        response = client.post(
            f"/api/articles/update/{test_article.id}",
            json={
                "nom": test_article.nom,
                "designation": test_article.designation,
                "reference": test_article.reference,
                "prix_achat_HT": 100.0,
                "location_price": 150.0,
                "taux_tva": 20.0,
            },
            headers=auth_headers,
        )

        assert response.status_code == 200

        # Verify location_price was updated
        response = client.get(
            f"/api/articles/info/{test_article.id}", headers=auth_headers
        )
        assert response.status_code == 200
        article_data = response.get_json()
        assert article_data["location_price"] == 150.0

    def test_get_article_info_includes_location_price(
        self, client, auth_headers, test_article
    ):
        """Test that article info response includes location_price field."""
        response = client.get(
            f"/api/articles/info/{test_article.id}", headers=auth_headers
        )

        assert response.status_code == 200
        data = response.get_json()
        assert "location_price" in data


class TestArticlesDelete:
    """Tests for deleting articles."""

    def test_delete_article_requires_auth(self, client, test_article):
        """Test that deleting an article requires authentication."""
        response = client.delete(f"/api/articles/delete/{test_article.id}")

        assert response.status_code == 401

    def test_delete_article_success(self, client, auth_headers, db_session):
        """Test successful article deletion."""
        from models import Articles, TauxTVA

        # Create a TVA rate
        tva = TauxTVA.query.filter_by(taux=20.0).first()
        if not tva:
            tva = TauxTVA(taux=20.0)
            db_session.add(tva)
            db_session.commit()

        # Create an article to delete
        article = Articles(
            nom="Article to Delete",
            designation="DEL-001",
            reference="REF-DEL-001",
            prix_achat_HT=100.0,
            prix_vente_HT=120.0,
            taux_tva_id=tva.id,
        )
        db_session.add(article)
        db_session.commit()
        article_id = article.id

        response = client.delete(
            f"/api/articles/delete/{article_id}", headers=auth_headers
        )

        assert response.status_code == 200
        data = response.get_json()
        assert data["message"] == "Article supprimé avec succès"

    def test_delete_nonexistent_article(self, client, auth_headers):
        """Test deleting non-existent article."""
        response = client.delete("/api/articles/delete/99999", headers=auth_headers)

        assert response.status_code == 404

    def test_delete_article_used_in_devis_returns_400(
        self, client, auth_headers, test_article, test_devis
    ):
        """Test deleting an article that is already used in a devis."""
        response = client.delete(
            f"/api/articles/delete/{test_article.id}", headers=auth_headers
        )

        assert response.status_code == 400
        assert "ne peut pas être supprimé" in response.get_json()["error"]


def _build_import_xlsx(rows):
    """Build an in-memory .xlsx matching the expected import format."""
    workbook = Workbook()
    worksheet = workbook.active

    # Rows 1-5 are ignored by the import logic.
    for index in range(1, 6):
        worksheet.cell(row=index, column=1, value=f"header-{index}")

    for index, row in enumerate(rows, start=6):
        worksheet.cell(row=index, column=2, value=row.get("reference"))
        worksheet.cell(row=index, column=3, value=row.get("nom"))
        worksheet.cell(row=index, column=4, value=row.get("prix"))

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


class TestArticlesExportPDF:
    """Tests for the articles PDF export endpoint."""

    def test_export_articles_pdf_requires_auth(self, client):
        response = client.get("/api/articles/export-pdf")
        assert response.status_code == 401

    def test_export_articles_pdf_no_articles_returns_404(self, client, auth_headers):
        response = client.get("/api/articles/export-pdf", headers=auth_headers)
        assert response.status_code == 404

    def test_export_articles_pdf_success(
        self, client, auth_headers, test_article, db_session, monkeypatch
    ):
        """Ensure PDF export returns a downloadable PDF payload."""

        rendered_context = {}

        class FakeHTML:
            def __init__(self, string, base_url):
                self.string = string
                self.base_url = base_url

            def write_pdf(self):
                return b"%PDF-1.4 fake"

        def fake_render_template(*args, **kwargs):
            rendered_context.update(kwargs)
            return "html"

        monkeypatch.setattr("routes.articles.render_template", fake_render_template)
        monkeypatch.setattr("routes.articles.HTML", FakeHTML)

        response = client.get("/api/articles/export-pdf", headers=auth_headers)

        assert response.status_code == 200
        assert response.headers["Content-Type"] == "application/pdf"
        assert (
            "attachment; filename=liste_articles_"
            in response.headers["Content-Disposition"]
        )
        assert "articles" in rendered_context
        assert (
            rendered_context["articles"][0]["location_price"]
            == test_article.location_price
        )

    def test_export_articles_pdf_exception_returns_500(
        self, client, auth_headers, test_article, monkeypatch
    ):
        def _raise_template_error(*args, **kwargs):
            raise RuntimeError("template failure")

        monkeypatch.setattr("routes.articles.render_template", _raise_template_error)

        response = client.get("/api/articles/export-pdf", headers=auth_headers)

        assert response.status_code == 500
        assert "Erreur lors de la génération du PDF" in response.get_json()["error"]


class TestArticlesImportXLSX:
    """Tests for the articles XLSX import endpoint."""

    def _multipart_headers(self, auth_headers):
        return {k: v for k, v in auth_headers.items() if k.lower() != "content-type"}

    def test_import_articles_xlsx_requires_auth(self, client):
        response = client.post("/api/articles/import-xlsx")
        assert response.status_code == 401

    def test_import_articles_xlsx_missing_file(self, client, auth_headers):
        response = client.post(
            "/api/articles/import-xlsx", headers=self._multipart_headers(auth_headers)
        )
        assert response.status_code == 400
        assert "Fichier Excel manquant" in response.get_json()["error"]

    def test_import_articles_xlsx_invalid_extension(self, client, auth_headers):
        response = client.post(
            "/api/articles/import-xlsx",
            data={"file": (io.BytesIO(b"invalid"), "articles.csv")},
            headers=self._multipart_headers(auth_headers),
        )
        assert response.status_code == 400
        assert "Format de fichier invalide" in response.get_json()["error"]

    def test_import_articles_xlsx_unreadable_file(self, client, auth_headers):
        response = client.post(
            "/api/articles/import-xlsx",
            data={"file": (io.BytesIO(b"not an xlsx"), "articles.xlsx")},
            headers=self._multipart_headers(auth_headers),
        )
        assert response.status_code == 400
        assert "illisible" in response.get_json()["error"]

    def test_import_articles_xlsx_incomplete_file(self, client, auth_headers):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.cell(row=1, column=1, value="header")

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = client.post(
            "/api/articles/import-xlsx",
            data={"file": (buffer, "articles.xlsx")},
            headers=self._multipart_headers(auth_headers),
        )
        assert response.status_code == 400
        assert "incomplet" in response.get_json()["error"]

    def test_import_articles_xlsx_no_tva_available(
        self, client, auth_headers, db_session
    ):
        from models import TauxTVA

        db_session.query(TauxTVA).delete()
        db_session.commit()

        buffer = _build_import_xlsx(
            [{"reference": "NEW-REF-001", "nom": "Article 1", "prix": 120}]
        )
        response = client.post(
            "/api/articles/import-xlsx",
            data={"file": (buffer, "articles.xlsx")},
            headers=self._multipart_headers(auth_headers),
        )

        assert response.status_code == 400
        assert "Aucun taux de TVA" in response.get_json()["error"]

    def test_import_articles_xlsx_mixed_results(
        self, client, auth_headers, db_session, test_article
    ):
        from models import Articles

        rows = [
            {"reference": "NEW-REF-001", "nom": "Nouvel article", "prix": "100,5"},
            {"reference": "NEW-REF-001", "nom": "Doublon", "prix": 140},
            {"reference": test_article.reference, "nom": "Existant", "prix": 200},
            {"reference": None, "nom": "No ref", "prix": 100},
            {"reference": "NO-NAME", "nom": None, "prix": 100},
            {"reference": "BAD-PRICE", "nom": "Prix invalide", "prix": "abc"},
        ]
        buffer = _build_import_xlsx(rows)

        response = client.post(
            "/api/articles/import-xlsx",
            data={"file": (buffer, "articles.xlsx")},
            headers=self._multipart_headers(auth_headers),
        )

        assert response.status_code == 200
        payload = response.get_json()
        assert payload["imported"] == 1
        assert len(payload["skipped"]) == 2
        assert len(payload["errors"]) == 3

        imported_article = Articles.query.filter_by(reference="NEW-REF-001").first()
        assert imported_article is not None
        assert imported_article.designation == ""
