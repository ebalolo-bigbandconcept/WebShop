import io
import logging
import os
import re
import unicodedata
from datetime import datetime

from flask import Blueprint, jsonify, make_response, render_template, request, session
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from weasyprint import HTML

from models import Articles, ArticlesSchema, DevisArticles, Parameters, TauxTVA, db
from utils import validate_article_fields

from .admin import admin_required

# Create a Blueprint for articles-related routes
articles_bp = Blueprint("articles_bp", __name__, url_prefix="/api/articles")


### Articles routes ###
# Get all articles info route
@articles_bp.route("/all", methods=["GET"])
def get_all_articles():
    # Check if user is authenticated
    user_id = session.get("user_id")
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    # Get all articles without pagination (frontend handles its own pagination)
    articles = Articles.query.order_by(Articles.id.asc()).all()

    if not articles:
        return jsonify({"error": "Aucuns articles trouvé"}), 404

    articles_schema = ArticlesSchema(many=True)
    articles_data = articles_schema.dump(articles)

    return jsonify({"data": articles_data})


# Get specific article info route
@articles_bp.route("/info/<article_id>", methods=["GET"])
def get_article_info(article_id):
    # Check if user is authenticated
    user_id = session.get("user_id")
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    article = Articles.query.filter_by(id=article_id).first()
    if not article:
        return jsonify({"error": "Article non trouvé"}), 404

    article_schema = ArticlesSchema()
    return article_schema.jsonify(article)


# Add new article route
@articles_bp.route("/create", methods=["POST"])
def add_article():
    # Check if user is authenticated
    user_id = session.get("user_id")
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    payload = request.get_json(force=True) or {}
    nom = payload["nom"]
    designation = payload.get("designation")
    if designation is None:
        designation = payload.get("reference", "")
        reference = payload.get("reference_plain", "")
    else:
        reference = payload.get("reference", "")
    prix_achat_HT = payload["prix_achat_HT"]
    taux_tva = payload["taux_tva"]
    location_price = payload.get("location_price")

    if taux_tva is None:
        return jsonify({"error": "Le taux de TVA est requis."}), 400
    try:
        taux_val = float(taux_tva)
    except (TypeError, ValueError):
        return jsonify({"error": "Le taux de TVA est invalide."}), 400
    # Accept either 0.20 or 20.0 style inputs
    if taux_val > 1:
        taux_val = taux_val / 100.0
    tva_obj = TauxTVA.query.filter_by(taux=taux_val).first()
    if not tva_obj:
        return jsonify({"error": "Taux TVA introuvable"}), 400
    taux_tva_id = tva_obj.id

    # Normalize location_price if provided
    if location_price is not None:
        if isinstance(location_price, str):
            location_price = location_price.replace(",", ".")
        location_price = float(location_price)

    # Get margin rate from parameters and calculate selling price
    params = Parameters.query.first()
    margin_rate = params.margin_rate if params else 0.0
    prix_vente_HT = float(prix_achat_HT) * margin_rate

    # Default location_price to 0 if not provided
    if location_price is None:
        location_price = 0.0

    error = validate_article_fields(
        nom,
        designation,
        reference,
        prix_achat_HT,
        prix_vente_HT,
        taux_tva_id,
        location_price,
    )
    if error:
        return jsonify({"error": error}), 400

    new_article = Articles(
        nom=nom,
        designation=designation,
        reference=reference,
        prix_achat_HT=prix_achat_HT,
        prix_vente_HT=prix_vente_HT,
        location_price=location_price,
        taux_tva_id=taux_tva_id,
    )
    db.session.add(new_article)
    db.session.commit()
    logging.info(
        f"Nouvel article ajouté: {new_article.nom} (id: {new_article.id}) par l'utilisateur {session.get('user_id')}"
    )

    return jsonify({"id": new_article.id})


# Modify article route
@articles_bp.route("/update/<article_id>", methods=["POST"])
def modify_article(article_id):
    # Check if user is authenticated
    user_id = session.get("user_id")
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    article = Articles.query.filter_by(id=article_id).first()
    if not article:
        return jsonify({"error": "Article non trouvé"}), 404

    payload = request.get_json(force=True) or {}
    new_nom = payload["nom"]
    new_designation = payload.get("designation")
    if new_designation is None:
        new_designation = payload.get("reference", "")
        new_reference = payload.get("reference_plain", "")
    else:
        new_reference = payload.get("reference", "")
    new_prix_achat_HT = payload["prix_achat_HT"]
    new_taux_tva = payload["taux_tva"]
    new_location_price = payload.get("location_price")

    if new_taux_tva is None:
        return jsonify({"error": "Le taux de TVA est requis."}), 400
    try:
        new_taux_val = float(new_taux_tva)
    except (TypeError, ValueError):
        return jsonify({"error": "Le taux de TVA est invalide."}), 400
    if new_taux_val > 1:
        new_taux_val = new_taux_val / 100.0
    new_tva_obj = TauxTVA.query.filter_by(taux=new_taux_val).first()
    if not new_tva_obj:
        return jsonify({"error": "Taux TVA introuvable"}), 400
    new_taux_tva_id = new_tva_obj.id

    # Normalize location_price if provided
    if new_location_price is not None:
        if isinstance(new_location_price, str):
            new_location_price = new_location_price.replace(",", ".")
        new_location_price = float(new_location_price)

    params = Parameters.query.first()
    margin_rate = params.margin_rate if params else 0.0
    new_prix_vente_HT = float(new_prix_achat_HT) * margin_rate

    # Default location_price to 0 if not provided
    if new_location_price is None:
        new_location_price = 0.0

    error = validate_article_fields(
        new_nom,
        new_designation,
        new_reference,
        new_prix_achat_HT,
        new_prix_vente_HT,
        new_taux_tva_id,
        new_location_price,
    )
    if error:
        return jsonify({"error": error}), 400

    article.nom = new_nom
    article.designation = new_designation
    article.reference = new_reference
    article.prix_achat_HT = new_prix_achat_HT
    article.prix_vente_HT = new_prix_vente_HT
    article.location_price = new_location_price
    article.taux_tva_id = new_taux_tva_id

    db.session.commit()
    logging.info(
        f"Article modifié: {article.nom} (id: {article.id}) par l'utilisateur {session.get('user_id')}"
    )

    return jsonify({"id": article.id})


# Delete article route
@articles_bp.route("/delete/<article_id>", methods=["DELETE"])
def delete_article(article_id):
    # Check if user is authenticated
    user_id = session.get("user_id")
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    article = Articles.query.filter_by(id=article_id).first()
    if not article:
        return jsonify({"error": "Article non trouvé"}), 404

    usage_count = DevisArticles.query.filter_by(article_id=article_id).count()
    if usage_count > 0:
        return (
            jsonify(
                {
                    "error": "Cet article est utilisé dans un ou plusieurs devis et ne peut pas être supprimé."
                }
            ),
            400,
        )

    article_name = article.nom
    try:
        Articles.query.filter_by(id=article_id).delete()
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return (
            jsonify(
                {
                    "error": "Cet article est utilisé dans un ou plusieurs devis et ne peut pas être supprimé."
                }
            ),
            400,
        )

    logging.info(
        f"Article supprimé: {article_name} (id: {article_id}) par l'utilisateur {session.get('user_id')}"
    )

    return jsonify({"message": "Article supprimé avec succès"})


# Export all articles as PDF.
@articles_bp.route("/export-pdf", methods=["GET"])
def export_articles_pdf():
    # Check if user is authenticated
    user_id = session.get("user_id")
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    try:
        articles = Articles.query.order_by(Articles.id.asc()).all()

        if not articles:
            return jsonify({"error": "Aucuns articles trouvé"}), 404

        # Get company info from parameters
        params = Parameters.query.first()
        company_name = params.company_name if params else "Artech Sécurité"

        # Format current date
        current_date = datetime.now().strftime("%d/%m/%Y")

        # Prepare articles data
        articles_data = []
        for article in articles:
            articles_data.append(
                {
                    "id": article.id,
                    "nom": article.nom,
                    "designation": article.designation,
                    "reference": article.reference,
                    "prix_achat_HT": article.prix_achat_HT,
                    "prix_vente_HT": article.prix_vente_HT,
                    "location_price": article.location_price,
                    "taux_tva": article.taux_tva,
                }
            )

        # Render HTML using Jinja2 template
        html_out = render_template(
            "articles_list.html",
            articles=articles_data,
            date=current_date,
            company_name=company_name,
        )

        # Calculate the absolute path to the pdf folder
        base_path = "/app/pdf/"

        # Generate PDF
        pdf_bytes = HTML(string=html_out, base_url=base_path).write_pdf()

        # Return PDF as response
        response = make_response(pdf_bytes)
        response.headers["Content-Type"] = "application/pdf"
        response.headers["Content-Disposition"] = (
            f"attachment; filename=liste_articles_{current_date.replace('/', '-')}.pdf"
        )
        return response

    except Exception as e:
        logging.error(f"Error generating articles PDF: {e}")
        return jsonify({"error": "Erreur lors de la génération du PDF"}), 500


@articles_bp.route("/import-xlsx", methods=["POST"])
def import_articles_xlsx():
    user_id = session.get("user_id")
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    upload = request.files.get("file")
    if not upload or upload.filename == "":
        return jsonify({"error": "Fichier Excel manquant."}), 400

    filename = upload.filename.lower()
    if not filename.endswith(".xlsx"):
        return jsonify({"error": "Format de fichier invalide. Utilisez un .xlsx."}), 400

    try:
        workbook = load_workbook(io.BytesIO(upload.read()), data_only=True)
    except Exception:
        return jsonify({"error": "Fichier Excel illisible."}), 400

    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return jsonify({"error": "Fichier Excel vide."}), 400

    def parse_price(value):
        if value is None:
            raise ValueError("Prix d'achat HT manquant")
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            cleaned = (
                value.strip().replace(" ", "").replace("\u00a0", "").replace(",", ".")
            )
            if cleaned == "":
                raise ValueError("Prix d'achat HT manquant")
            return float(cleaned)
        raise ValueError("Prix d'achat HT invalide")

    if len(rows) < 6:
        return jsonify({"error": "Fichier Excel incomplet."}), 400

    # Fixed column positions (0-indexed):
    # Column B (index 1): Référence
    # Column C (index 2): Désignation (nom)
    # Column D (index 3): Prix d'achat HT
    COL_REFERENCE = 1
    COL_NOM = 2
    COL_PRIX = 3

    tva_obj = TauxTVA.query.filter_by(taux=0.20).first()
    if not tva_obj:
        tva_obj = TauxTVA.query.order_by(TauxTVA.id.asc()).first()
    if not tva_obj:
        return jsonify({"error": "Aucun taux de TVA disponible."}), 400

    params = Parameters.query.first()
    margin_rate = params.margin_rate if params else 0.0

    parsed_rows = []
    row_errors = []

    # Start reading from row 6 (index 5)
    for row_index, row in enumerate(rows[5:], start=6):
        if (
            row is None
            or len(row) <= COL_PRIX
            or all(cell in (None, "") for cell in row)
        ):
            continue

        reference_raw = row[COL_REFERENCE] if len(row) > COL_REFERENCE else None
        nom_raw = row[COL_NOM] if len(row) > COL_NOM else None
        prix_raw = row[COL_PRIX] if len(row) > COL_PRIX else None

        reference = str(reference_raw).strip() if reference_raw is not None else ""
        nom = str(nom_raw).strip() if nom_raw is not None else ""

        if not reference:
            row_errors.append({"row": row_index, "error": "Référence manquante"})
            continue
        if not nom:
            row_errors.append({"row": row_index, "error": "Nom de l'article manquant"})
            continue
        try:
            prix_achat_ht = parse_price(prix_raw)
        except (TypeError, ValueError) as exc:
            row_errors.append({"row": row_index, "error": str(exc)})
            continue

        parsed_rows.append(
            {
                "row": row_index,
                "reference": reference,
                "nom": nom,
                "prix_achat_HT": prix_achat_ht,
            }
        )

    existing_refs = set()
    incoming_refs = {item["reference"] for item in parsed_rows if item["reference"]}
    if incoming_refs:
        existing_refs = {
            item.reference
            for item in Articles.query.filter(
                Articles.reference.in_(incoming_refs)
            ).all()
            if item.reference
        }

    skipped = []
    imported = 0
    seen_refs = set()

    for item in parsed_rows:
        reference = item["reference"]
        if reference in seen_refs:
            skipped.append(
                {
                    "row": item["row"],
                    "reference": reference,
                    "reason": "Doublon dans le fichier",
                }
            )
            continue
        seen_refs.add(reference)

        if reference in existing_refs:
            skipped.append(
                {
                    "row": item["row"],
                    "reference": reference,
                    "reason": "Référence déjà existante",
                }
            )
            continue

        prix_vente_ht = float(item["prix_achat_HT"]) * margin_rate
        designation = ""

        error = validate_article_fields(
            item["nom"],
            designation,
            reference,
            item["prix_achat_HT"],
            prix_vente_ht,
            tva_obj.id,
            0.0,
        )
        if error:
            row_errors.append({"row": item["row"], "error": error})
            continue

        new_article = Articles(
            nom=item["nom"],
            designation=designation,
            reference=reference,
            prix_achat_HT=item["prix_achat_HT"],
            prix_vente_HT=prix_vente_ht,
            location_price=0.0,
            taux_tva_id=tva_obj.id,
        )
        db.session.add(new_article)
        imported += 1

    if imported:
        db.session.commit()

    logging.info(
        "Import articles Excel: %s importes, %s ignores, %s erreurs (utilisateur %s)",
        imported,
        len(skipped),
        len(row_errors),
        user_id,
    )

    return jsonify(
        {
            "imported": imported,
            "skipped": skipped,
            "errors": row_errors,
        }
    )
